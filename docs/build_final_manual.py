from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"F:\Users\kangc\Desktop\Lost-Found重构实施与验收手册.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="4"):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.text = text
        set_cell_shading(cell, "365F91")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cells[index], "F2F6FA")
            for paragraph in cells[index].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph()
    return table


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    return paragraph


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("校园失物招领系统重构实施与验收手册")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("若依管理平台 独立校园用户端 单体模块化后端")
    run.font.size = Pt(12)
    run.bold = True
    doc.add_paragraph("目标工程：G:\\VibeCodingFiles\\Lost-Found")
    doc.add_paragraph("原始项目：G:\\MInnngyue\\Lost-Found")
    doc.add_paragraph("文档版本：最终实施版，2026 年 9 月 6 日")
    doc.add_paragraph(
        "本手册记录重构后的工程结构、技术选型、业务规则、实施步骤、验证结果和上线前置条件。"
        "原始项目和原数据库均保留，重构工程使用独立的 ry-vue-lostfound 数据库。"
    )

    add_heading(doc, "1 结论与交付边界", 1)
    doc.add_paragraph(
        "本项目已完成从原始 Vibe Coding 项目到 RuoYi Vue 管理平台、独立校园用户端和 Spring Boot 4 "
        "模块化单体后端的核心重构。后端、管理端和门户均可构建；本地服务已启动并完成真实数据库接口验收。"
    )
    add_table(doc, ["项目项", "当前结果", "证据或说明"], [
        ["后端框架", "已完成", "RuoYi 3.9.2 体系，Spring Boot 4.1.0，Java 17 编译目标，Java 21 运行"],
        ["持久层", "已完成", "MyBatis，Mapper XML 参数绑定，分页和业务服务分层"],
        ["管理端", "已完成", "RuoYi Vue3，Vite，端口 http://localhost/"],
        ["校园用户端", "已完成", "Vue3，Vite，端口 http://localhost:5173/"],
        ["后端服务", "已完成", "端口 http://localhost:8081/，连接 MySQL 8 和 Redis"],
        ["数据库", "已完成", "独立数据库 ry-vue-lostfound，原 ry-vue 未修改"],
        ["备份恢复", "本地演练通过", "v2 备份结构、角色和上传文件已核对"],
        ["生产边界", "上线前置", "WAF/CDN/DDoS、HTTPS、最小权限账号和正式对象存储仍需部署"]
    ], [3.2, 3.0, 9.5])

    add_heading(doc, "2 总体架构", 1)
    doc.add_paragraph(
        "系统采用单体模块化后端，而不是把业务拆成多个微服务。这样可以在校园项目规模下保持部署简单、"
        "事务边界清晰和调试成本可控，同时通过 Maven 模块和业务服务边界保留后续演进空间。"
    )
    add_code(doc, "浏览器 -> portal 校园用户端 -> backend 单体后端 -> MySQL / Redis / 文件存储")
    add_code(doc, "浏览器 -> admin 若依管理端 -> backend 单体后端 -> RuoYi 系统模块 + ruoyi-lostfound")
    add_table(doc, ["模块", "主要职责", "部署和边界"], [
        ["portal", "广场、详情、发布、认领、评论、消息、私信、举报、申诉、个人中心", "独立 Vue3 应用，不直连数据库"],
        ["admin", "用户、角色、菜单、分类、审核、认领、治理、日志", "独立 RuoYi Vue3 应用"],
        ["ruoyi-admin", "应用入口、系统 Controller、登录和公共接口", "单体进程入口"],
        ["ruoyi-system", "用户、角色、菜单、字典、配置、日志", "复用 RuoYi 平台能力"],
        ["ruoyi-framework", "JWT、Spring Security、Redis、文件、异常、限流", "提供横切安全能力"],
        ["ruoyi-lostfound", "帖子、图片、分类、匹配、认领、互动、治理", "业务模块，规则集中在 Service"],
        ["legacy-backend", "旧实现对照和回滚参考", "不参与新系统运行"]
    ], [3.1, 7.2, 5.4])

    add_heading(doc, "3 技术栈与选择理由", 1)
    add_table(doc, ["层次", "技术", "选择理由"], [
        ["后端", "Spring Boot 4.1.0", "匹配 RuoYi Spring Boot 4 主线，统一 Jakarta 命名空间和现代依赖"],
        ["管理平台", "RuoYi Vue3 3.9.2", "复用成熟的用户、角色、菜单、按钮权限、日志、代码生成和定时任务"],
        ["认证授权", "Spring Security + JWT + Redis", "无状态接口认证，配合 RuoYi 方法级权限和登录缓存"],
        ["持久层", "MyBatis", "与 RuoYi 主线一致，Mapper SQL 可审查，减少同时维护两套持久层风格"],
        ["数据库", "MySQL 8", "承载 RuoYi 系统表和 lf_ 业务表，支持事务、索引和备份恢复"],
        ["缓存", "Redis", "验证码、登录状态、字典缓存、限流和后续未读消息缓存"],
        ["校园端", "Vue 3 + Vite + Element Plus + Pinia + Axios", "保留原端技术基础，集中适配 RuoYi API"],
        ["管理端", "Vue 3 + Vite + Element Plus + Pinia", "使用 RuoYi 动态路由、权限指令和后台组件"],
        ["测试", "JUnit 5 + Mockito + Node test", "覆盖服务状态机、Mapper XML、前端发布校验和关键边界"],
        ["构建", "Maven Wrapper + npm lockfile + GitHub Actions", "固定构建入口和依赖版本，支持持续集成"]
    ], [3.0, 5.0, 7.7])
    doc.add_paragraph(
        "本项目明确选择 MyBatis，不同时引入 MyBatis-Plus。除非未来出现明确的批量更新或复杂查询收益，"
        "否则不应为了代码简短而改变已经与 RuoYi 对齐的持久层风格。"
    )

    add_heading(doc, "4 业务重构结果", 1)
    add_heading(doc, "4.1 帖子和审核", 2)
    for text in [
        "帖子类型为 LOST 寻物和 FOUND 招领，类型不再和状态混用。",
        "生命周期为 DRAFT、PENDING_REVIEW、PUBLISHED、RESOLVED、EXPIRED、WITHDRAWN、REJECTED、REMOVED。",
        "只有审核通过且已发布的帖子进入公开广场；草稿、待审核、拒绝、下架、过期和已解决内容不公开。",
        "标题、帖子类型、物品分类、颜色、校区、事件日期等必填字段由前端提示并由后端再次校验。",
        "发布、编辑、提交审核、审核通过、拒绝、撤回、过期和解决均写入状态事件或审核记录。",
        "拒绝和下架必须保留原因，用户可以编辑被拒绝帖子并重新提交，不产生重复帖子。"
    ]:
        add_bullet(doc, text)
    add_heading(doc, "4.2 分类和地点", 2)
    doc.add_paragraph(
        "统一使用 lf_category 管理 ITEM、COLOR、CAMPUS、AREA、LOCATION 五类数据。业务层拒绝非法类型、"
        "跨校区区域、错误地点父级、停用父级和循环引用；存在子节点时不得随意改变父级、类型或停用父级。"
    )
    add_heading(doc, "4.3 图片", 2)
    for text in [
        "最多 9 张图片，单张不超过 5 MB。",
        "按真实文件内容识别图片格式，不信任用户提交的扩展名或 MIME 字段。",
        "使用 SHA-256 去重，文件名不直接使用用户原始文件名。",
        "图片访问根据帖子公开状态、用户归属和审核员权限判断；文件路径经过规范化并限制在上传目录内。",
        "数据库保存 object_key、摘要、大小、实际 MIME 类型和图片审核状态，上传文件单独备份。"
    ]:
        add_bullet(doc, text)
    add_heading(doc, "4.4 匹配和认领", 2)
    doc.add_paragraph(
        "匹配只表示系统建议，不代表所有权或认领已经成立。匹配按物品分类、颜色、地点和日期排序，"
        "认领则通过申请、接受、拒绝、取消、交接、双方确认和完成形成独立状态机。重复请求通过业务校验、"
        "事务和版本控制保持幂等；认领完成后关联帖子进入 RESOLVED。"
    )
    add_heading(doc, "4.5 信用分替代方案", 2)
    doc.add_paragraph(
        "已删除公开信用分、信用等级和按分数限制业务的设计。新的惩罚机制是可解释、可追溯的治理链："
        "举报 -> 违规案件 -> 处罚动作 -> 功能限制 -> 申诉 -> 审计事件。处罚范围可以是发布、认领、私信或全部功能，"
        "并记录原因、证据、处理人和时间。"
    )
    add_heading(doc, "4.6 互动与治理", 2)
    add_table(doc, ["能力", "实现内容"], [
        ["评论", "帖子评论、长度校验、权限校验、举报入口和内容状态"],
        ["系统通知", "审核结果、匹配、认领、评论、举报和申诉结果通知"],
        ["私信", "按会话查询，只允许发送者或接收者访问，具备频率限制和举报入口"],
        ["举报", "帖子、评论、用户、认领过程和私信均可举报，支持处理状态"],
        ["违规案件", "OPEN、INVESTIGATING、RESOLVED、CLOSED，支持案件处理和审计"],
        ["处罚", "警告、发布限制、认领限制、私信限制、下架、临时封禁、永久封禁"],
        ["申诉", "被限制用户提交申诉，运营管理员审核并保留结果"],
    ], [4.0, 11.7])

    add_heading(doc, "5 数据库和迁移策略", 1)
    doc.add_paragraph(
        "RuoYi 平台表使用 sys_ 前缀，失物招领业务表使用 lf_ 前缀。业务表不通过物理外键强耦合 sys_user，"
        "而由 Service 校验用户存在、状态和对象归属。这样既保留 RuoYi 用户体系，也避免业务迁移被平台表结构锁死。"
    )
    add_table(doc, ["类别", "表"], [
        ["帖子", "lf_post、lf_post_image、lf_post_review、lf_post_event"],
        ["匹配和认领", "lf_match、lf_claim、lf_claim_event"],
        ["互动", "lf_comment、lf_message、lf_private_message"],
        ["治理", "lf_report、lf_violation_case、lf_violation_action、lf_appeal、lf_audit_event"],
        ["分类和版本", "lf_category、lf_schema_version"],
        ["RuoYi 平台", "sys_user、sys_role、sys_menu、sys_user_role、sys_role_menu 等"]
    ], [5.0, 10.7])
    doc.add_paragraph("重要 SQL 文件：")
    for text in [
        "backend/sql/ry_20260417.sql：RuoYi 系统表和基础数据。",
        "backend/sql/lost_found_v2_fresh.sql：新业务表结构，只用于空业务库。",
        "backend/sql/upgrade_ry_vue_lostfound_to_v2.sql：从业务 v1 升级到 v2。",
        "backend/sql/lost_found_category_seed.sql：补齐默认分类，不重置用户自建分类。",
        "backend/sql/lost_found_role_seed.sql：审核员和运营管理员角色及权限。",
        "backend/sql/lost_found_common_role_grant.sql：收敛 common 角色到门户所需权限。",
        "backend/sql/lost_found_restore_verify.sql：只读恢复校验，不应对活动库执行破坏性操作。"
    ]:
        add_bullet(doc, text)

    add_heading(doc, "6 安全基线", 1)
    add_table(doc, ["风险", "当前工程控制", "仍需的生产控制"], [
        ["SQL 注入", "MyBatis 参数绑定，Druid Wall 禁止多语句", "依赖升级、最小权限数据库账号、数据库网络隔离"],
        ["越权访问", "Spring Security 方法权限 + Service 对象归属校验", "持续权限回归和生产账号审计"],
        ["暴力和滥用", "Redis 限流覆盖登录相关业务、发布、上传、评论、私信、举报", "WAF、网关、IP/账号/设备风险策略"],
        ["文件攻击", "真实图片识别、大小和数量限制、SHA-256、路径规范化", "对象存储隔离、病毒扫描、配额和生命周期策略"],
        ["XSS", "RuoYi XSS 过滤和内容长度约束", "富文本白名单、CSP 和安全响应头"],
        ["敏感配置", "密码、JWT、Redis 和上传路径使用环境变量", "密钥管理系统和轮换机制"],
        ["DDoS", "应用层限流和受限 CORS 只能降低部分滥用", "CDN/WAF/云 DDoS 防护、源站隐藏和应急扩容"]
    ], [3.0, 7.0, 5.7])
    doc.add_paragraph(
        "必须明确：本地 Spring Boot 限流不是生产级 DDoS 防护。正式上线必须由 CDN、云 WAF、负载均衡或云厂商 "
        "DDoS 防护承接公网流量，MySQL 和 Redis 不能直接暴露公网。"
    )

    add_heading(doc, "7 从零部署和运行步骤", 1)
    add_heading(doc, "7.1 环境准备", 2)
    for text in [
        "Java 17 或更高版本，当前本机使用 Java 21 运行。",
        "Node.js 20.19 或更高版本。",
        "MySQL 8，Redis 或 Memurai。",
        "Navicat 仅用于初始化、迁移和备份恢复，不作为应用运行依赖。",
        "不要修改原始项目 G:\\MInnngyue\\Lost-Found 和原数据库 ry-vue。"
    ]:
        add_number(doc, text)
    add_heading(doc, "7.2 初始化数据库", 2)
    for text in [
        "创建 ry-vue-lostfound 数据库，并导入 RuoYi 基础 SQL。",
        "导入 backend/sql/lost_found_v2_fresh.sql。",
        "依次导入分类、角色、菜单、认领匹配、互动和治理种子 SQL。",
        "对已有 v1 业务库，先完整备份，再执行 v1 到 v2 升级 SQL。",
        "在 Navicat 中核对 sys_user、sys_role、sys_menu、sys_role_menu 和全部 lf_ 表。"
    ]:
        add_number(doc, text)
    add_heading(doc, "7.3 启动后端", 2)
    add_code(doc, "cd G:\\VibeCodingFiles\\Lost-Found")
    add_code(doc, "powershell.exe -NoProfile -ExecutionPolicy Bypass -File .\\scripts\\Start-LostFound.ps1")
    doc.add_paragraph(
        "脚本以掩码方式读取 MySQL 密码，并只在当前 Java 进程期间注入 LOSTFOUND_DB_PASSWORD。"
        "也可以使用 scripts\\Start-LostFound.cmd，适用于 PowerShell 禁止执行 ps1 的环境。"
    )
    add_heading(doc, "7.4 启动前端", 2)
    add_code(doc, "cd G:\\VibeCodingFiles\\Lost-Found\\admin")
    add_code(doc, "npm.cmd run dev")
    add_code(doc, "cd G:\\VibeCodingFiles\\Lost-Found\\portal")
    add_code(doc, "npm.cmd run dev")
    doc.add_paragraph("当前本地地址：管理端 http://localhost/，门户 http://localhost:5173/，后端 http://localhost:8081/")

    add_heading(doc, "8 验收结果", 1)
    add_table(doc, ["验收项", "结果", "实际证据"], [
        ["后端测试", "通过", "51 tests，0 failures，0 errors"],
        ["门户测试", "通过", "12 tests，0 failures"],
        ["门户生产构建", "通过", "npm run build"],
        ["管理端生产构建", "通过", "npm run build:prod"],
        ["验证码和登录", "通过", "captchaImage 正常，admin 真实登录成功"],
        ["普通用户权限", "通过", "我的帖子、通知允许；系统用户、分类、管理帖子均 403"],
        ["审核员权限", "通过", "帖子和匹配管理允许；系统用户、分类、治理均 403"],
        ["运营管理员权限", "通过", "分类、认领、举报、案件允许；系统用户和角色均 403"],
        ["未登录保护", "通过", "受保护接口返回 RuoYi code 401"],
        ["公开帖子", "通过", "公开列表和详情可读，帖子图片返回真实 PNG"],
        ["分类前置条件", "通过", "ITEM 8，COLOR 11，CAMPUS 1 个启用根节点"],
        ["备份恢复结构", "通过", "v2 备份含 21 张预期表、4 个角色和图片记录"],
        ["恢复上传文件", "通过", "活动路径和 restore-test 副本均为 81,131 字节"]
    ], [4.2, 2.3, 9.2])
    doc.add_paragraph(
        "角色验收使用短用户名的临时测试账号，以避免 RuoYi 用户名长度规则干扰测试。测试完成后账号已删除。"
        "这次实测还发现并修复了 common 角色继承基础 RuoYi 系统菜单的问题，说明权限验收不是仅依赖前端菜单隐藏。"
    )

    add_heading(doc, "9 备份恢复操作规范", 1)
    for text in [
        "停止写入流量或至少停止业务操作。",
        "在 Navicat 中完整备份 ry-vue-lostfound 的结构和数据。",
        "恢复到独立数据库，例如 ry-vue-lostfound-restore-test-v2，不覆盖活动库。",
        "执行 backend/sql/lost_found_restore_verify.sql，只读核对数据库名称、表、行数、版本、角色、分类、帖子和图片。",
        "按 ruoyi.profile 备份并恢复上传目录，核对 lf_post_image.object_key 对应文件。",
        "用恢复配置启动第二个后端实例，完成一次登录、帖子详情和图片访问冒烟。",
        "通过后再把恢复库交给正式切换或灾难恢复流程；禁止直接拿原库做恢复测试。"
    ]:
        add_number(doc, text)
    doc.add_paragraph("已生成的本地备份：_data\\backups\\ry-vue-lostfound_full_2026-09-06-v2.sql")

    add_heading(doc, "10 后续生产上线清单", 1)
    for text in [
        "创建只拥有 ry-vue-lostfound 必要权限的应用数据库账号，禁止应用使用 root。",
        "配置生产 JWT 密钥、数据库密码、Redis 密码和上传路径到密钥管理系统。",
        "使用 HTTPS、HSTS、可信 CORS 来源和安全响应头。",
        "在应用前部署 CDN、WAF、DDoS 防护和反向代理，隐藏源站地址。",
        "将 MySQL、Redis 和上传存储放入内网或受控网络。",
        "关闭生产 SQL 参数明文日志和不必要的 Druid、Swagger 入口。",
        "将本地文件存储迁移到隔离的对象存储，并增加病毒扫描和配额。",
        "执行登录、注册、发布、上传、评论、认领、私信和举报的压力与滥用测试。",
        "配置进程自动重启、健康检查、日志轮转、监控和告警。",
        "完成独立恢复环境的登录、帖子详情和图片访问冒烟测试。"
    ]:
        add_bullet(doc, text)

    add_heading(doc, "11 项目文件索引", 1)
    add_table(doc, ["路径", "用途"], [
        ["backend", "RuoYi Spring Boot 4 模块化单体后端"],
        ["admin", "RuoYi Vue3 管理端"],
        ["portal", "独立校园用户端"],
        ["legacy-backend", "原后端迁移参考和回滚依据"],
        ["docs/REFACTOR_EXECUTION.md", "阶段状态、决策和执行证据"],
        ["docs/ACCEPTANCE_RUNBOOK.md", "人工和自动验收步骤"],
        ["docs/BUSINESS_REQUIREMENTS_V1.md", "业务边界和状态流转规范"],
        ["docs/DATA_MODEL_V1.md", "数据模型和表职责"],
        ["docs/SECURITY_BASELINE.md", "安全基线和生产边界"],
        ["docs/DEPLOYMENT.md", "本地和生产部署说明"],
        ["scripts/Start-LostFound.cmd", "绕过 PowerShell 策略的安全启动入口"]
    ], [7.0, 8.7])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Lost-Found 重构实施与验收手册")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
